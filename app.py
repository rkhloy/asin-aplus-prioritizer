import re
from io import BytesIO
from collections import Counter
from typing import Dict, List

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


st.set_page_config(page_title="ASIN A+ Prioritizer", layout="wide")

st.title("ASIN A+ Prioritizer")
st.caption(
    "Upload listing data, score SKU quality, and export a clean prioritization file."
)

uploaded_file = st.file_uploader("Upload CSV or Excel file", type=["csv", "xlsx"])


if "analysis_ready" not in st.session_state:
    st.session_state.analysis_ready = False
if "output_df" not in st.session_state:
    st.session_state.output_df = None
if "source_file_name" not in st.session_state:
    st.session_state.source_file_name = ""
if "last_uploaded_name" not in st.session_state:
    st.session_state.last_uploaded_name = ""
if "current_mapping" not in st.session_state:
    st.session_state.current_mapping = None


def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def safe_number(value):
    try:
        if pd.isna(value):
            return 0.0
        text = str(value).strip().replace("$", "").replace(",", "")
        if text == "":
            return 0.0
        return float(text)
    except Exception:
        return 0.0


def looks_like_url(value):
    text = clean_text(value).lower()
    return (
        text.startswith("http://")
        or text.startswith("https://")
        or text.startswith("www.")
    )


def get_text_value(row, column_name):
    if not column_name or column_name == "-- Not Mapped --":
        return ""
    return clean_text(row.get(column_name, ""))


def get_value(row, column_name):
    if not column_name or column_name == "-- Not Mapped --":
        return None
    return row.get(column_name, None)


def suggest_column(columns, keywords):
    for col in columns:
        col_clean = str(col).strip().lower()
        for keyword in keywords:
            if keyword in col_clean:
                return col
    return "-- Not Mapped --"


def get_index(options, suggested):
    if suggested in options:
        return options.index(suggested)
    return 0


def add_reason(reasons, category_counter, reason, category):
    reasons.append(reason)
    category_counter[category] += 1


def add_tip(tips: List[str], text: str):
    if text not in tips:
        tips.append(text)


def format_excel_export(output_buffer, export_df):
    with pd.ExcelWriter(output_buffer, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="Results")
        worksheet = writer.sheets["Results"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        header_fill = PatternFill(fill_type="solid", fgColor="1F2937")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        for col_idx, col_name in enumerate(export_df.columns, start=1):
            max_length = len(str(col_name))
            sample_values = export_df[col_name].astype(str).fillna("").head(50).tolist()
            for val in sample_values:
                max_length = max(max_length, len(val))
            worksheet.column_dimensions[get_column_letter(col_idx)].width = min(
                max(max_length + 2, 12), 36
            )


def analyze_row(row, col_map):
    reasons = []
    tips = []
    categories = Counter()
    critical_issue = False

    sku = get_text_value(row, col_map["sku"])
    asin = get_text_value(row, col_map["asin"])
    title = get_text_value(row, col_map["title"])
    brand = get_text_value(row, col_map["brand"])
    category = get_text_value(row, col_map["category"])
    description = get_text_value(row, col_map["description"])
    meta_title = get_text_value(row, col_map["meta_title"])
    channel = get_text_value(row, col_map["channel"])

    bullet_values = [
        get_text_value(row, col_map[key])
        for key in ["bullet1", "bullet2", "bullet3", "bullet4"]
    ]
    feature_values = [
        get_text_value(row, col_map[key])
        for key in ["feature1", "feature2", "feature3", "feature4"]
    ]
    image_values = [
        get_text_value(row, col_map[key])
        for key in ["image1", "image2", "image3", "image4", "image5"]
    ]

    default_price = safe_number(get_value(row, col_map["default_price"]))
    sale_price = safe_number(get_value(row, col_map["sale_price"]))

    populated_bullets = [x for x in bullet_values if x]
    populated_features = [x for x in feature_values if x]
    populated_images = [x for x in image_values if x]

    content_score = 0
    if len(title) >= 60:
        content_score += 10
    elif len(title) >= 30:
        content_score += 7
    elif len(title) >= 10:
        content_score += 3
    else:
        add_reason(reasons, categories, "Missing or weak title", "Content")
        add_tip(tips, "Improve product title")
        critical_issue = True

    if len(description) >= 120:
        content_score += 10
    elif len(description) >= 40:
        content_score += 6
    else:
        add_reason(reasons, categories, "Missing or weak description", "Content")
        add_tip(tips, "Improve product description")
        if len(description) == 0:
            critical_issue = True

    bullet_feature_count = len(populated_bullets) + len(populated_features)
    if bullet_feature_count >= 6:
        content_score += 10
    elif bullet_feature_count >= 3:
        content_score += 6
    elif bullet_feature_count >= 1:
        content_score += 3
    else:
        add_reason(reasons, categories, "Missing bullets and features", "Content")
        add_tip(tips, "Add bullets or feature fields")

    image_score = 0
    if image_values[0] and looks_like_url(image_values[0]):
        image_score += 10
    else:
        add_reason(reasons, categories, "Primary image missing or invalid", "Images")
        add_tip(tips, "Add valid primary image URL")
        critical_issue = True

    valid_extra_images = sum(1 for x in image_values[1:] if x and looks_like_url(x))
    if valid_extra_images >= 4:
        image_score += 10
    elif valid_extra_images >= 2:
        image_score += 6
    elif valid_extra_images >= 1:
        image_score += 3
    else:
        add_reason(reasons, categories, "Low image coverage", "Images")
        add_tip(tips, "Add more gallery images")

    pricing_score = 0
    if default_price > 0:
        pricing_score += 10
    else:
        add_reason(reasons, categories, "Invalid default price", "Pricing")
        add_tip(tips, "Fix default price")
        critical_issue = True

    if col_map["sale_price"] != "-- Not Mapped --":
        if sale_price == 0 or sale_price <= default_price:
            pricing_score += 10
        else:
            add_reason(
                reasons, categories, "Sale price exceeds default price", "Pricing"
            )
            add_tip(tips, "Review sale price logic")
            critical_issue = True
    else:
        pricing_score += 10

    metadata_score = 0
    if brand:
        metadata_score += 5
    else:
        add_reason(reasons, categories, "Missing brand", "Metadata")
        add_tip(tips, "Populate brand")

    if category:
        metadata_score += 5
    else:
        add_reason(reasons, categories, "Missing category", "Metadata")
        add_tip(tips, "Populate category")

    if meta_title:
        metadata_score += 5
    else:
        add_reason(reasons, categories, "Missing meta title", "Metadata")
        add_tip(tips, "Add meta title")

    if channel:
        metadata_score += 5

    total_score = round(content_score + image_score + pricing_score + metadata_score, 1)

    if critical_issue or total_score < 60:
        priority = "High"
    elif total_score < 80:
        priority = "Medium"
    else:
        priority = "Low"

    priority_rank = 1 if priority == "High" else 2 if priority == "Medium" else 3
    top_issue_category = categories.most_common(1)[0][0] if categories else "Healthy"
    failure_reasons = "; ".join(dict.fromkeys(reasons))
    optimization_tips = "; ".join(dict.fromkeys(tips))
    smart_a_plus_readiness = "Yes" if total_score >= 80 and not critical_issue else "No"

    if priority == "High":
        recommended_next_action = (
            tips[0] if tips else "Fix highest-impact listing issues"
        )
    elif priority == "Medium":
        recommended_next_action = tips[0] if tips else "Address warnings"
    else:
        recommended_next_action = "Monitor listing quality"

    return {
        "optimization_score": total_score,
        "priority": priority,
        "priority_rank": priority_rank,
        "top_issue_category": top_issue_category,
        "smart_a_plus_readiness": smart_a_plus_readiness,
        "issue_count": len(set(reasons)),
        "failure_reasons": failure_reasons,
        "optimization_tips": optimization_tips,
        "recommended_next_action": recommended_next_action,
        "sku_value": sku,
        "asin_value": asin,
        "title_value": title,
    }


if uploaded_file is not None:
    try:
        file_name = uploaded_file.name.lower()
        if file_name.endswith(".csv"):
            df = pd.read_csv(uploaded_file)
        elif file_name.endswith(".xlsx"):
            df = pd.read_excel(uploaded_file)
        else:
            st.error("Unsupported file type. Please upload a CSV or Excel file.")
            st.stop()

        df.columns = [str(col).strip() for col in df.columns]
        raw_columns = df.columns.tolist()
        columns = ["-- Not Mapped --"] + raw_columns

        st.sidebar.header("Column Mapping")

        default_map = {
            "sku": suggest_column(raw_columns, ["sku", "item sku"]),
            "asin": suggest_column(raw_columns, ["asin"]),
            "title": suggest_column(
                raw_columns,
                ["title", "product title", "product name", "name", "item name"],
            ),
            "brand": suggest_column(raw_columns, ["brand"]),
            "category": suggest_column(raw_columns, ["category", "product type"]),
            "description": suggest_column(raw_columns, ["description"]),
            "meta_title": suggest_column(
                raw_columns, ["meta title", "meta", "seo title"]
            ),
            "channel": suggest_column(
                raw_columns, ["channel", "source", "marketplace"]
            ),
            "bullet1": suggest_column(raw_columns, ["bullet1", "bullet 1"]),
            "bullet2": suggest_column(raw_columns, ["bullet2", "bullet 2"]),
            "bullet3": suggest_column(raw_columns, ["bullet3", "bullet 3"]),
            "bullet4": suggest_column(raw_columns, ["bullet4", "bullet 4"]),
            "feature1": suggest_column(raw_columns, ["feature1", "feature 1"]),
            "feature2": suggest_column(raw_columns, ["feature2", "feature 2"]),
            "feature3": suggest_column(raw_columns, ["feature3", "feature 3"]),
            "feature4": suggest_column(raw_columns, ["feature4", "feature 4"]),
            "image1": suggest_column(raw_columns, ["image1", "image 1", "main image"]),
            "image2": suggest_column(raw_columns, ["image2", "image 2"]),
            "image3": suggest_column(raw_columns, ["image3", "image 3"]),
            "image4": suggest_column(raw_columns, ["image4", "image 4"]),
            "image5": suggest_column(raw_columns, ["image5", "image 5"]),
            "default_price": suggest_column(
                raw_columns, ["default price", "price", "regular price"]
            ),
            "sale_price": suggest_column(
                raw_columns, ["sale price", "special price", "promo price"]
            ),
        }

        required_fields = [
            "sku",
            "title",
            "brand",
            "category",
            "description",
            "image1",
            "default_price",
        ]
        optional_fields = [
            "asin",
            "meta_title",
            "channel",
            "bullet1",
            "bullet2",
            "bullet3",
            "bullet4",
            "feature1",
            "feature2",
            "feature3",
            "feature4",
            "image2",
            "image3",
            "image4",
            "image5",
            "sale_price",
        ]

        col_map: Dict[str, str] = {}
        with st.sidebar.expander("Required fields", expanded=True):
            for field in required_fields:
                col_map[field] = st.selectbox(
                    f"{field.replace('_', ' ').title()} *",
                    columns,
                    index=get_index(columns, default_map[field]),
                    key=f"map_{field}",
                )

        with st.sidebar.expander("Optional fields", expanded=False):
            for field in optional_fields:
                col_map[field] = st.selectbox(
                    field.replace("_", " ").title(),
                    columns,
                    index=get_index(columns, default_map[field]),
                    key=f"map_{field}",
                )

        st.markdown(
            '### Selected Mapping <span style="color:#ef4444;">*required</span>',
            unsafe_allow_html=True,
        )
        mapping_df = pd.DataFrame(
            {
                "Field": list(col_map.keys()),
                "Selected Column": list(col_map.values()),
                "Type": [
                    "Required" if f in required_fields else "Optional"
                    for f in col_map.keys()
                ],
            }
        )
        st.dataframe(mapping_df, use_container_width=True, hide_index=True)

        missing_required = [
            field for field in required_fields if col_map[field] == "-- Not Mapped --"
        ]
        if missing_required:
            st.error("Missing required mappings: " + ", ".join(missing_required))
            st.stop()

        current_uploaded_name = uploaded_file.name if uploaded_file is not None else ""
        if current_uploaded_name != st.session_state.last_uploaded_name:
            st.session_state.analysis_ready = False
            st.session_state.output_df = None
            st.session_state.source_file_name = ""
            st.session_state.last_uploaded_name = current_uploaded_name

        action_col1, action_col2 = st.columns([3.6, 1])
        run_analysis = action_col1.button(
            "Run Analysis", type="primary", use_container_width=True
        )
        clear_clicked = action_col2.button("Clear", use_container_width=True)

        if clear_clicked:
            st.session_state.analysis_ready = False
            st.session_state.output_df = None
            st.session_state.source_file_name = ""
            st.rerun()

        if run_analysis:
            results_df = df.apply(
                lambda row: pd.Series(analyze_row(row, col_map)), axis=1
            )
            output_df = pd.concat([df.copy(), results_df], axis=1)

            front_cols = [
                col_map["sku"],
                col_map["asin"],
                col_map["title"],
                col_map["brand"],
                col_map["category"],
                col_map["channel"],
                "optimization_score",
                "priority",
                "priority_rank",
                "top_issue_category",
                "smart_a_plus_readiness",
                "issue_count",
                "recommended_next_action",
                "failure_reasons",
                "optimization_tips",
            ]
            front_cols = [
                c
                for c in front_cols
                if c in output_df.columns and c != "-- Not Mapped --"
            ]
            remaining_cols = [c for c in output_df.columns if c not in front_cols]
            output_df = output_df[front_cols + remaining_cols]

            st.session_state.output_df = output_df
            st.session_state.analysis_ready = True
            st.session_state.source_file_name = file_name
            st.session_state.current_mapping = col_map.copy()

        if st.session_state.analysis_ready and st.session_state.output_df is not None:
            output_df = st.session_state.output_df.copy()
            active_map = (
                st.session_state.current_mapping
                if st.session_state.current_mapping
                else col_map
            )

            total_rows = len(output_df)
            high_count = int((output_df["priority"] == "High").sum())
            medium_count = int((output_df["priority"] == "Medium").sum())
            low_count = int((output_df["priority"] == "Low").sum())
            avg_score = (
                round(output_df["optimization_score"].mean(), 1) if total_rows else 0
            )

            st.subheader("Summary")
            summary_df = pd.DataFrame(
                [[total_rows, high_count, medium_count, low_count, avg_score]],
                columns=[
                    "Total Rows",
                    "High Priority",
                    "Medium Priority",
                    "Low Priority",
                    "Average Score",
                ],
            )
            st.dataframe(summary_df, use_container_width=True, hide_index=True)

            export_file_name = st.session_state.source_file_name.lower()
            if export_file_name.endswith(".csv"):
                output_data = output_df.to_csv(index=False).encode("utf-8")
                output_file_name = "asin_aplus_prioritizer_results.csv"
                output_mime = "text/csv"
            else:
                output_buffer = BytesIO()
                format_excel_export(output_buffer, output_df)
                output_data = output_buffer.getvalue()
                output_file_name = "asin_aplus_prioritizer_results.xlsx"
                output_mime = (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

            st.download_button(
                "Download Results File",
                data=output_data,
                file_name=output_file_name,
                mime=output_mime,
                use_container_width=True,
            )

            st.subheader("Results")
            sort_by = st.selectbox(
                "Sort By",
                options=["Priority Rank", "Optimization Score", "Issue Count"],
                index=0,
            )

            filtered_df = output_df.copy()
            if sort_by == "Priority Rank":
                filtered_df = filtered_df.sort_values(
                    by=["priority_rank", "optimization_score"], ascending=[True, True]
                )
            elif sort_by == "Optimization Score":
                filtered_df = filtered_df.sort_values(
                    by=["optimization_score"], ascending=[True]
                )
            else:
                filtered_df = filtered_df.sort_values(
                    by=["issue_count", "optimization_score"], ascending=[False, True]
                )

            st.dataframe(filtered_df, use_container_width=True)

    except Exception as e:
        st.error(f"Error reading file: {e}")
